"""Offline source fidelity contracts. Estimated/actual OpenRouter cost: $0/$0."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import fitz
import pytest
from docx import Document
from docx.oxml import parse_xml

from app.services.docx_extract import extract_docx
from app.services.document_processor import DocumentProcessor, _extraction_failure_details
from app.services.extraction_quality import OCRResult, assess_pdf_page, extract_source
from app.services.retrieval import select_evidence
from app.services.vector_store import VectorStore


OMML = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _math(xml: str):
    return parse_xml(f'<m:oMath xmlns:m="{OMML}">{xml}</m:oMath>')


def test_docx_preserves_body_order_nested_tables_and_math(tmp_path: Path):
    path = tmp_path / "math.docx"
    doc = Document()
    doc.add_paragraph("Opening paragraph")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Critical value: alpha = 42"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested evidence"
    p = doc.add_paragraph("Equation: ")
    p._p.append(_math('<m:r><m:t>x+y=2</m:t></m:r>'))
    frac = doc.add_paragraph("Fraction: ")
    frac._p.append(_math(
        '<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num>'
        '<m:den><m:r><m:t>b</m:t></m:r></m:den></m:f>'
    ))
    doc.save(path)

    blocks, report = extract_docx(path)

    assert [block.kind for block in blocks] == ["paragraph", "table", "paragraph", "math", "paragraph", "math"]
    assert blocks[0].text == "Opening paragraph"
    assert "alpha = 42" in blocks[1].text
    assert "Nested evidence" in blocks[1].text
    assert blocks[3].math_latex == "x+y=2"
    assert blocks[5].math_latex == r"\frac{a}{b}"
    assert blocks[1].location.startswith("body[1]/table")
    assert all(block.page is None for block in blocks)
    assert report.complete


def test_unknown_omml_is_retained_and_marks_incomplete(tmp_path: Path):
    path = tmp_path / "unknown.docx"
    doc = Document()
    p = doc.add_paragraph("Unsupported: ")
    p._p.append(_math('<m:phant><m:e><m:r><m:t>visible</m:t></m:r></m:e></m:phant>'))
    doc.save(path)

    blocks, report = extract_docx(path)

    assert any("visible" in (block.text or block.math_latex or "") for block in blocks)
    assert not report.complete
    assert any("unsupported OMML" in warning and "body[0]" in warning for warning in report.warnings)


def test_inline_text_math_text_has_unique_locations_and_live_chunk_ids(tmp_path: Path):
    path = tmp_path / "inline.docx"
    doc = Document()
    paragraph = doc.add_paragraph("Before equation ")
    paragraph._p.append(_math('<m:r><m:t>x=1</m:t></m:r>'))
    paragraph.add_run(" after equation with enough source words to remain meaningful")
    doc.save(path)

    blocks, _ = extract_docx(path)
    chunks = DocumentProcessor(vector_store=None, chunk_size=1000).extract_and_chunk_file(str(path), "course")

    assert [block.text or block.math_latex for block in blocks] == [
        "Before equation", "x=1", "after equation with enough source words to remain meaningful"
    ]
    assert len({block.location for block in blocks}) == 3
    assert len({chunk.metadata["chunk_id"] for chunk in chunks}) == 3
    captured = {}
    vector_store = VectorStore.__new__(VectorStore)
    vector_store.collection_name = "identity-contract"
    vector_store.collection = type(
        "RecordingCollection", (), {"upsert": lambda _self, **kwargs: captured.update(kwargs)}
    )()
    vector_store.embedding_model = "deterministic-test"
    vector_store.embedding_dimensions = 2
    vector_store.normalization_version = "v1"
    vector_store._embedding_service = type(
        "DeterministicEmbeddings",
        (),
        {"embed_texts": lambda _self, texts, **_kwargs: [[0.0, 1.0] for _ in texts]},
    )()
    vector_store.add_documents(chunks, course_id="course")
    assert len(captured["ids"]) == len(set(captured["ids"])) == 3
    assert captured["documents"] == [chunk.content for chunk in chunks]


def test_omml_scripts_brace_compound_and_nested_bases(tmp_path: Path):
    path = tmp_path / "scripts.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph._p.append(_math(
        '<m:sSup><m:e><m:r><m:t>a+b</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>'
        '<m:sSubSup><m:e><m:sSub><m:e><m:r><m:t>x+y</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub></m:e>'
        '<m:sub><m:r><m:t>j</m:t></m:r></m:sub><m:sup><m:r><m:t>3</m:t></m:r></m:sup></m:sSubSup>'
    ))
    doc.save(path)

    blocks, report = extract_docx(path)

    assert blocks[0].math_latex == "{a+b}^{2}{{x+y}_{i}}_{j}^{3}"
    assert report.complete


def _scan_pdf(path: Path, count: int, *, native_first: bool = False):
    source = fitz.open()
    source_page = source.new_page()
    source_page.insert_text((72, 72), "Rendered scan content")
    pix = source_page.get_pixmap()
    png = pix.tobytes("png")
    source.close()
    out = fitz.open()
    if native_first:
        page = out.new_page()
        page.insert_textbox(fitz.Rect(50, 50, 500, 200), "Native text page with enough readable content for extraction and indexing.")
    for _ in range(count):
        page = out.new_page()
        page.insert_image(page.rect, stream=png)
    out.save(path)
    out.close()


def test_pdf_assesses_each_page_and_ocr_cap_counts_omissions(tmp_path: Path):
    path = tmp_path / "mixed.pdf"
    _scan_pdf(path, 13, native_first=True)
    calls = []

    def ocr(_image: bytes, page: int):
        calls.append(page)
        return OCRResult(text=f"OCR page {page}", complete=True, latency_ms=page, input_tokens=0, output_tokens=0)

    blocks, report = extract_source(path, ocr_page=ocr, ocr_max_pages=12)

    assert calls == list(range(2, 14))
    assert report.total_pages == 14
    assert report.ocr_pages == list(range(2, 14))
    assert report.skipped_pages == [14]
    assert report.extracted_pages == list(range(1, 14))
    assert not report.complete
    assert {block.page for block in blocks} == set(range(1, 14))


def test_truncated_ocr_keeps_partial_text_and_marks_damaged(tmp_path: Path):
    path = tmp_path / "scan.pdf"
    _scan_pdf(path, 1)

    blocks, report = extract_source(
        path,
        ocr_page=lambda _image, _page: OCRResult(
            text="partial transcription", complete=False, finish_reason="length", latency_ms=4
        ),
    )

    assert blocks[0].text == "partial transcription"
    assert report.ocr_pages == [1]
    assert report.ocr_details[0]["latency_ms"] == 4
    assert report.ocr_details[0]["finish_reason"] == "length"
    assert report.extracted_pages == [1]
    assert report.damaged_pages == [1]
    assert not report.complete
    assert any("truncated" in warning for warning in report.warnings)


def test_valid_empty_ocr_marks_image_only_page_uncovered(tmp_path: Path):
    path = tmp_path / "empty-ocr.pdf"
    _scan_pdf(path, 1)

    blocks, report = extract_source(
        path, ocr_page=lambda _image, _page: OCRResult(text="", complete=True, finish_reason="stop")
    )

    assert blocks == []
    assert report.ocr_pages == [1]
    assert report.extracted_pages == []
    assert report.damaged_pages == [1]
    assert not report.complete
    assert any("no usable text" in warning for warning in report.warnings)


def test_confirmed_blank_page_is_counted_without_ocr_or_damage(tmp_path: Path):
    path = tmp_path / "blank.pdf"
    pdf = fitz.open()
    pdf.new_page()
    pdf.save(path)
    pdf.close()
    calls = []

    blocks, report = extract_source(
        path, ocr_page=lambda _image, page: calls.append(page) or OCRResult(text="", complete=True)
    )

    assert blocks == []
    assert calls == []
    assert report.blank_pages == [1]
    assert report.damaged_pages == []
    assert report.complete


def test_inline_pdf_image_is_not_certified_blank_when_ocr_is_empty(tmp_path: Path):
    path = tmp_path / "inline-image.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    xref = pdf.get_new_xref()
    pdf.update_object(xref, "<<>>")
    pdf.update_stream(
        xref,
        b"q 200 0 0 200 72 500 cm BI /W 2 /H 2 /CS /RGB /BPC 8 /F /AHx "
        b"ID FF000000FF000000FFFFFFFF> EI Q",
    )
    page.set_contents(xref)
    pdf.save(path)
    pdf.close()
    calls = []

    blocks, report = extract_source(
        path,
        ocr_page=lambda _image, page_number: calls.append(page_number)
        or OCRResult(text="", complete=True, finish_reason="stop"),
    )

    assert blocks == []
    assert calls == [1]
    assert report.blank_pages == []
    assert report.ocr_pages == [1]
    assert report.damaged_pages == [1]
    assert not report.complete


def test_real_damaged_symbol_assessment_selects_ocr_and_retains_unrepaired_damage(tmp_path: Path):
    path = tmp_path / "damaged.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Native page represented by many damaged glyphs")
    pdf.save(path)
    pdf.close()
    damaged_text = "\ue000" * 80 + " retained native words"
    assessment = assess_pdf_page(
        text=damaged_text,
        spans=[{"text": damaged_text, "bbox": (0, 0, 100, 12)}],
        image_count=0,
        min_chars=50,
    )
    assert assessment.needs_ocr
    assert assessment.damaged
    calls = []

    def ocr(_image, page_number):
        calls.append(page_number)
        return OCRResult(text="", complete=True, latency_ms=1)

    blocks, report = extract_source(
        path,
        ocr_page=ocr,
        page_analysis_inputs={
            1: {"text": damaged_text, "spans": [{"text": damaged_text, "bbox": (0, 0, 100, 12)}], "image_count": 0}
        },
    )

    assert calls == [1]
    assert "Native page represented" in blocks[0].text
    assert "/bbox[" in blocks[0].location
    assert report.ocr_pages == [1]
    assert report.damaged_pages == [1]
    assert not report.complete


def test_real_assessor_detects_low_ratio_replacement_glyph_damage():
    damaged_text = "readable " * 70 + "\ue000" * 10

    assessment = assess_pdf_page(
        text=damaged_text,
        spans=[{"text": damaged_text, "bbox": (0, 0, 500, 12)}],
        image_count=0,
        drawing_count=0,
        min_chars=50,
    )

    assert assessment.damaged
    assert assessment.needs_ocr
    assert not assessment.confirmed_blank


def test_canonical_chunking_restores_headers_front_matter_and_keeps_short_structures(tmp_path: Path):
    pdf_path = tmp_path / "noise.pdf"
    pdf = fitz.open()
    bodies = [
        "ISBN 978-1-234 Publisher Copyright Edition details for this source document.",
        "A substantive chapter explains neural systems with enough words for grounded retrieval and study.",
        "Another substantive chapter explains optimization methods with enough words for grounded retrieval and study.",
    ]
    for index, body in enumerate(bodies, 1):
        page = pdf.new_page()
        page.insert_text((72, 72), f"REPEATED HEADER\n{body}\nREPEATED FOOTER")
    pdf.save(pdf_path)
    pdf.close()
    processor = DocumentProcessor(vector_store=None, chunk_size=1000)

    chunks = processor.extract_and_chunk_file(str(pdf_path), "course")

    assert all("REPEATED HEADER" not in chunk.content for chunk in chunks)
    assert all("REPEATED FOOTER" not in chunk.content for chunk in chunks)
    assert any(chunk.metadata["is_front_matter"] is True for chunk in chunks)
    assert any(chunk.metadata["is_front_matter"] is False for chunk in chunks)

    docx_path = tmp_path / "short-structures.docx"
    doc = Document()
    doc.add_paragraph("x")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "42"
    math = doc.add_paragraph()
    math._p.append(_math('<m:r><m:t>x</m:t></m:r>'))
    doc.save(docx_path)
    structured = processor.extract_and_chunk_file(str(docx_path), "course")
    assert {chunk.metadata["block_kind"] for chunk in structured} == {"table", "math"}


def test_txt_bom_is_supported_but_invalid_utf8_is_rejected(tmp_path: Path):
    good = tmp_path / "good.txt"
    good.write_bytes(b"\xef\xbb\xbfStrict UTF-8 text")
    blocks, report = extract_source(good)
    assert blocks[0].text == "Strict UTF-8 text"
    assert report.complete

    bad = tmp_path / "bad.txt"
    bad.write_bytes(b"bad\xfftext")
    with pytest.raises(UnicodeDecodeError):
        extract_source(bad)
    try:
        bad.read_text(encoding="utf-8", errors="strict")
    except UnicodeDecodeError as exc:
        message, code, action = _extraction_failure_details(exc)
    assert "UTF-8" in message
    assert code == "DOCUMENT_TEXT_ENCODING_UNSUPPORTED"
    assert action == "resave_utf8"


@pytest.mark.parametrize("suffix,payload", [(".pdf", b"%PDF-1.7\ninvalid"), (".docx", b"not-a-zip")])
def test_malformed_binary_documents_have_no_text_fallback(tmp_path: Path, suffix: str, payload: bytes):
    path = tmp_path / f"bad{suffix}"
    path.write_bytes(payload)
    with pytest.raises(Exception):
        extract_source(path, ocr_page=lambda _image, _page: OCRResult(text="", complete=False))


def test_malformed_docx_package_is_rejected(tmp_path: Path):
    path = tmp_path / "missing-document-part.docx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
    with pytest.raises(Exception):
        extract_docx(path)


def test_canonical_blocks_reach_chunk_metadata(tmp_path: Path):
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("A sufficiently detailed paragraph containing source context for the canonical indexing path.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "A sufficiently detailed table containing alpha = 42 and supporting evidence."
    doc.save(path)

    chunks = DocumentProcessor(vector_store=None, chunk_size=1000).extract_and_chunk_file(str(path), "course")

    assert {chunk.metadata["block_kind"] for chunk in chunks} == {"paragraph", "table"}
    assert all("document_location" in chunk.metadata for chunk in chunks)
    assert all("page" not in chunk.metadata for chunk in chunks)


def test_actual_ingestion_metadata_preserves_canonical_chunk_presentation_order(tmp_path: Path):
    path = tmp_path / "ordered.docx"
    doc = Document()
    doc.add_paragraph(" ".join(f"First canonical sentence {index}." for index in range(12)))
    doc.add_paragraph(" ".join(f"Second canonical sentence {index}." for index in range(12)))
    doc.save(path)
    chunks = DocumentProcessor(
        vector_store=None,
        chunk_size=100,
        chunk_overlap=10,
    ).extract_and_chunk_file(str(path), "course")

    assert all("block_order" in chunk.metadata for chunk in chunks)
    assert all("within_block_order" in chunk.metadata for chunk in chunks)
    selected = select_evidence(list(reversed(chunks)), k=len(chunks))
    assert [chunk.metadata["chunk_id"] for chunk in selected] == [
        chunk.metadata["chunk_id"] for chunk in chunks
    ]
